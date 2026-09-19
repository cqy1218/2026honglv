import csv
import html
import json
import os
import re
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook


# ============================================================
# 基础配置
# ============================================================

EXCEL_FILE = Path("参考逻辑.xlsx")

DATA_DIR = Path("data")
JSON_FILE = DATA_DIR / "scores.json"
CSV_FILE = DATA_DIR / "scores.csv"

README_FILE = Path("README.md")


# ============================================================
# 真实网页对应关系
#
# 以网页/API作为唯一的：
# - 分组依据
# - 排名依据
# - 得分依据
# - 项目名称依据
#
# sort_type=2 = 按得分高低顺序
# ============================================================

GROUPS = [
    {
        "name": "公益1组",
        "road_show_id": 150,
    },
    {
        "name": "公益2组",
        "road_show_id": 151,
    },
    {
        "name": "创意1组",
        "road_show_id": 152,
    },
    {
        "name": "创意2组",
        "road_show_id": 153,
    },
    {
        "name": "创意3组",
        "road_show_id": 154,
    },
]


BASE_URL = (
    "https://itcjspapif.woczx.com/"
    "contest-lives/road-show/{road_show_id}?sort_type=2"
)


# ============================================================
# HTTP 请求头
# ============================================================

HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Content-Type": "application/json",
    "Origin": "https://jspitcm.woczx.com",
    "Referer": "https://jspitcm.woczx.com/",
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-site",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/152.0.0.0 Safari/537.36"
    ),
    "X-Client-Type": "mobile",
    "sec-ch-ua-mobile": "?0",
    "sec-ch-ua-platform": '"Windows"',
}


# ============================================================
# 可选身份验证
# ============================================================

cookie = os.getenv(
    "WOCZX_COOKIE",
    "",
).strip()

authorization = os.getenv(
    "WOCZX_AUTHORIZATION",
    "",
).strip()

if cookie:
    HEADERS["Cookie"] = cookie

if authorization:
    HEADERS["Authorization"] = authorization


# ============================================================
# README 自动区域
# ============================================================

README_START = "<!-- SCOREBOARD_START -->"
README_END = "<!-- SCOREBOARD_END -->"


# ============================================================
# 通用函数
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_project_name(value):
    """
    用于 Excel 学校匹配。

    这里只做项目名称的容错处理，不参与网页分组。

    处理：
    - 全角/半角
    - μ / µ 等 Unicode 兼容字符
    - 空格
    - 中英文标点
    - 各种破折号
    - 引号

    例如：

    星驭——面向卫星互联网...
    星驭-面向卫星互联网...

    会得到比较接近的标准形式。
    """

    text = clean_text(value)

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFKC",
        text,
    )

    text = text.lower()

    # 删除空格、标点、下划线等
    text = re.sub(
        r"[\W_]+",
        "",
        text,
        flags=re.UNICODE,
    )

    return text


def rank_number(value):
    try:
        return int(value)

    except (
        TypeError,
        ValueError,
    ):
        return 999999


def format_rank(rank):
    if rank in (
        None,
        "",
    ):
        return ""

    return f"第{rank}名"


def escape(value):
    if value is None:
        return ""

    return html.escape(
        str(value)
    )


# ============================================================
# Excel 只负责：
#
# 项目名称 → 学校名称
#
# 不再使用：
# - 路演项目组
# - 答辩顺序
# - Excel中的排名
# - Excel中的组别
#
# ============================================================

def load_school_lookup():

    exact_lookup = {}

    normalized_candidates = {}

    if not EXCEL_FILE.exists():

        print()
        print(
            f"⚠️ 未找到 {EXCEL_FILE}"
        )

        print(
            "学校信息将显示为“未匹配”，"
            "但不影响网页成绩、项目和排名。"
        )

        return (
            exact_lookup,
            {},
        )

    print("=" * 70)
    print(
        f"读取学校信息：{EXCEL_FILE}"
    )
    print("=" * 70)

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    sheet_name = "路演答辩顺序"

    if sheet_name not in workbook.sheetnames:

        workbook.close()

        print(
            f"⚠️ Excel 中不存在“{sheet_name}”工作表。"
        )

        return (
            exact_lookup,
            {},
        )

    sheet = workbook[
        sheet_name
    ]

    # ========================================================
    # 自动查找标题行，而不是完全依赖固定行号
    # ========================================================

    header_row = None

    project_col = None
    school_col = None

    max_search_rows = min(
        sheet.max_row,
        10,
    )

    for row_number in range(
        1,
        max_search_rows + 1,
    ):

        values = [
            clean_text(
                cell.value
            )
            for cell in sheet[
                row_number
            ]
        ]

        for index, value in enumerate(
            values
        ):

            if value == "项目名称":
                project_col = index

            if value == "学校名称":
                school_col = index

        if (
            project_col is not None
            and school_col is not None
        ):

            header_row = row_number
            break

    # ========================================================
    # 如果自动识别失败，兼容原表：
    #
    # A = 项目名称
    # B = 学校名称
    # 第二行为标题
    # ========================================================

    if header_row is None:

        print(
            "⚠️ 未自动识别标题，"
            "按 A列项目名称 / B列学校名称读取。"
        )

        header_row = 2
        project_col = 0
        school_col = 1

    # ========================================================
    # 读取项目和学校
    # ========================================================

    count = 0

    for row in sheet.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    ):

        if not row:
            continue

        if len(row) <= max(
            project_col,
            school_col,
        ):
            continue

        project_name = clean_text(
            row[
                project_col
            ]
        )

        school_name = clean_text(
            row[
                school_col
            ]
        )

        if not project_name:
            continue

        if not school_name:
            continue

        count += 1

        # ----------------------------------------------------
        # 第一层：原名称精确匹配
        # ----------------------------------------------------

        if project_name not in exact_lookup:

            exact_lookup[
                project_name
            ] = school_name

        # ----------------------------------------------------
        # 第二层：标准化名称匹配
        # ----------------------------------------------------

        normalized = normalize_project_name(
            project_name
        )

        if not normalized:
            continue

        if normalized not in normalized_candidates:

            normalized_candidates[
                normalized
            ] = []

        normalized_candidates[
            normalized
        ].append(
            {
                "project":
                    project_name,

                "school":
                    school_name,
            }
        )

    workbook.close()

    # ========================================================
    # 只保留“没有歧义”的标准化匹配
    #
    # 如果同一个标准化名称对应多个学校，
    # 宁可不匹配，也绝不猜学校。
    # ========================================================

    normalized_lookup = {}

    for (
        normalized,
        candidates,
    ) in normalized_candidates.items():

        schools = {
            candidate[
                "school"
            ]
            for candidate in candidates
        }

        if len(
            schools
        ) == 1:

            normalized_lookup[
                normalized
            ] = candidates[
                0
            ]

    print(
        f"Excel 中读取到 {count} 条项目学校关系"
    )

    print(
        f"精确项目名：{len(exact_lookup)} 条"
    )

    print(
        f"可安全模糊兼容：{len(normalized_lookup)} 条"
    )

    print()

    return (
        exact_lookup,
        normalized_lookup,
    )


# ============================================================
# 根据真实网页项目名寻找学校
# ============================================================

def find_school(
    project_name,
    exact_lookup,
    normalized_lookup,
):

    # ========================================================
    # 第一优先级：
    # 完全一致
    # ========================================================

    if project_name in exact_lookup:

        return (
            exact_lookup[
                project_name
            ],
            "exact",
            "",
        )

    # ========================================================
    # 第二优先级：
    # 去标点后的唯一对应
    # ========================================================

    normalized = normalize_project_name(
        project_name
    )

    candidate = normalized_lookup.get(
        normalized
    )

    if candidate:

        return (
            candidate[
                "school"
            ],
            "normalized",
            candidate[
                "project"
            ],
        )

    # ========================================================
    # 找不到就明确显示未匹配
    #
    # 不再根据答辩顺序猜学校。
    # ========================================================

    return (
        "⚠️ 未匹配",
        "unmatched",
        "",
    )


# ============================================================
# 获取一个真实网页的数据
# ============================================================

def fetch_group(
    session,
    group,
):

    group_name = group[
        "name"
    ]

    road_show_id = group[
        "road_show_id"
    ]

    url = BASE_URL.format(
        road_show_id=road_show_id
    )

    print("=" * 70)

    print(
        f"读取：{group_name}"
    )

    print(
        f"road-show ID：{road_show_id}"
    )

    print(
        f"API：{url}"
    )

    response = session.get(
        url,
        timeout=30,
    )

    print(
        f"HTTP：{response.status_code}"
    )

    if response.status_code == 401:

        raise RuntimeError(
            f"{group_name} 返回 401 Unauthorized。"
        )

    response.raise_for_status()

    try:

        data = response.json()

    except Exception:

        print(
            "返回结果不是有效 JSON："
        )

        print(
            response.text[
                :1000
            ]
        )

        raise

    # ========================================================
    # 输出网页自己报告的小组名称
    # ========================================================

    contest_road_show = (
        data.get(
            "contestRoadShow",
            {},
        )
        or {}
    )

    judge_group = (
        contest_road_show.get(
            "judgeGroup",
            {},
        )
        or {}
    )

    api_group_name = clean_text(
        judge_group.get(
            "name"
        )
    )

    if api_group_name:

        print(
            f"网页实际小组名称：{api_group_name}"
        )

    return data


# ============================================================
# 解析一个网页
#
# 重要：
#
# 项目
# 得分
# 排名
# 路演顺序
#
# 全部直接以 API 为准。
# ============================================================

def parse_group(
    group,
    data,
    exact_lookup,
    normalized_lookup,
):

    group_name = group[
        "name"
    ]

    road_show_id = group[
        "road_show_id"
    ]

    projects = (
        data.get(
            "apply_project",
            [],
        )
        or []
    )

    result = []

    normalized_matches = []

    unmatched_projects = []

    for item in projects:

        project = (
            item.get(
                "project",
                {},
            )
            or {}
        )

        project_name = clean_text(
            project.get(
                "name"
            )
        )

        # ====================================================
        # 得分
        # ====================================================

        score = item.get(
            "grade",
            "",
        )

        # ====================================================
        # 排名
        #
        # 首选 rank
        # 其次 group_rank
        # ====================================================

        rank = item.get(
            "rank"
        )

        if rank in (
            None,
            "",
        ):

            rank = item.get(
                "group_rank"
            )

        # ====================================================
        # 路演顺序
        #
        # 只作为展示/调试信息。
        # 不拿它匹配学校。
        # ====================================================

        road_show_sort = item.get(
            "road_show_sort",
            "",
        )

        # ====================================================
        # 学校匹配
        # ====================================================

        (
            school,
            match_type,
            excel_project_name,
        ) = find_school(
            project_name,
            exact_lookup,
            normalized_lookup,
        )

        if match_type == "normalized":

            normalized_matches.append(
                {
                    "api":
                        project_name,

                    "excel":
                        excel_project_name,

                    "school":
                        school,
                }
            )

        elif match_type == "unmatched":

            unmatched_projects.append(
                {
                    "project":
                        project_name,

                    "road_show_sort":
                        road_show_sort,

                    "school_id":
                        project.get(
                            "school_id",
                            "",
                        ),
                }
            )

        result.append(
            {
                "group":
                    group_name,

                "road_show_id":
                    road_show_id,

                "school":
                    school,

                "school_match":
                    match_type,

                "project":
                    project_name,

                "score":
                    score,

                "rank":
                    rank,

                "rank_text":
                    format_rank(
                        rank
                    ),

                "road_show_sort":
                    road_show_sort,

                "project_id":
                    item.get(
                        "project_id",
                        "",
                    ),

                "school_id":
                    project.get(
                        "school_id",
                        "",
                    ),
            }
        )

    # ========================================================
    # 再按 API 排名排序
    # ========================================================

    result.sort(
        key=lambda x:
        rank_number(
            x[
                "rank"
            ]
        )
    )

    print(
        f"{group_name}：共 {len(result)} 个项目"
    )

    # ========================================================
    # 仅做学校名称检查
    # ========================================================

    if normalized_matches:

        print()

        print(
            f"ℹ️ {group_name} 有 "
            f"{len(normalized_matches)} 个项目"
            "通过名称标准化匹配学校："
        )

        for item in normalized_matches:

            print(
                f"   API：{item['api']}"
            )

            print(
                f"   Excel：{item['excel']}"
            )

            print(
                f"   学校：{item['school']}"
            )

    if unmatched_projects:

        print()

        print(
            f"⚠️ {group_name} 有 "
            f"{len(unmatched_projects)} 个项目"
            "没有匹配到学校："
        )

        for item in unmatched_projects:

            print(
                f"   {item['project']}"
            )

            print(
                f"      路演顺序："
                f"{item['road_show_sort']}"
            )

            print(
                f"      API school_id："
                f"{item['school_id']}"
            )

    print()

    return result


# ============================================================
# 用于比较本次榜单和上次榜单
# ============================================================

def comparable_data(
    groups_data,
):

    return {
        group_name: [
            {
                "road_show_id":
                    row[
                        "road_show_id"
                    ],

                "school":
                    row[
                        "school"
                    ],

                "project":
                    row[
                        "project"
                    ],

                "score":
                    row[
                        "score"
                    ],

                "rank":
                    row[
                        "rank"
                    ],

                "road_show_sort":
                    row[
                        "road_show_sort"
                    ],

                "project_id":
                    row[
                        "project_id"
                    ],

                "school_id":
                    row[
                        "school_id"
                    ],
            }

            for row in rows
        ]

        for (
            group_name,
            rows,
        ) in groups_data.items()
    }


# ============================================================
# 读取旧数据
# ============================================================

def load_previous_data():

    if not JSON_FILE.exists():

        return None

    try:

        with JSON_FILE.open(
            "r",
            encoding="utf-8",
        ) as f:

            data = json.load(
                f
            )

        if isinstance(
            data,
            dict,
        ):

            return data.get(
                "groups"
            )

    except Exception as exc:

        print(
            f"⚠️ 旧 scores.json 读取失败：{exc}"
        )

    return None


# ============================================================
# CSV
# ============================================================

def write_csv(
    groups_data,
):

    DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    fieldnames = [
        "组别",
        "road_show_id",
        "学校",
        "项目名称",
        "得分",
        "名次",
        "路演顺序",
        "project_id",
        "school_id",
        "学校匹配方式",
    ]

    with CSV_FILE.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()

        for (
            group_name,
            rows,
        ) in groups_data.items():

            for row in rows:

                writer.writerow(
                    {
                        "组别":
                            group_name,

                        "road_show_id":
                            row[
                                "road_show_id"
                            ],

                        "学校":
                            row[
                                "school"
                            ],

                        "项目名称":
                            row[
                                "project"
                            ],

                        "得分":
                            row[
                                "score"
                            ],

                        "名次":
                            row[
                                "rank_text"
                            ],

                        "路演顺序":
                            row[
                                "road_show_sort"
                            ],

                        "project_id":
                            row[
                                "project_id"
                            ],

                        "school_id":
                            row[
                                "school_id"
                            ],

                        "学校匹配方式":
                            row[
                                "school_match"
                            ],
                    }
                )


# ============================================================
# README 首页榜单
# ============================================================

def generate_scoreboard_html(
    groups_data,
    updated_at,
):

    group_names = [
        group[
            "name"
        ]
        for group in GROUPS
    ]

    max_rows = max(
        len(
            groups_data.get(
                group_name,
                [],
            )
        )
        for group_name in group_names
    )

    lines = []

    lines.append(
        '<div align="center">'
    )

    lines.append(
        "<h2>“建行杯”江苏大学生创新大赛（2026）实时成绩</h2>"
    )

    lines.append(
        f"<p>最近一次数据变化：{escape(updated_at)}</p>"
    )

    lines.append(
        "</div>"
    )

    lines.append("")

    lines.append(
        "<table>"
    )

    # ========================================================
    # 一级表头
    # ========================================================

    lines.append(
        "<thead>"
    )

    lines.append(
        "<tr>"
    )

    for group_name in group_names:

        lines.append(
            '<th colspan="4" align="center">'
            f"<strong>{escape(group_name)}</strong>"
            "</th>"
        )

    lines.append(
        "</tr>"
    )

    # ========================================================
    # 二级表头
    # ========================================================

    lines.append(
        "<tr>"
    )

    for _ in group_names:

        lines.append(
            '<th align="center">学校</th>'
        )

        lines.append(
            '<th align="center">项目名称</th>'
        )

        lines.append(
            '<th align="center">得分</th>'
        )

        lines.append(
            '<th align="center">名次</th>'
        )

    lines.append(
        "</tr>"
    )

    lines.append(
        "</thead>"
    )

    # ========================================================
    # 内容
    # ========================================================

    lines.append(
        "<tbody>"
    )

    for index in range(
        max_rows
    ):

        lines.append(
            "<tr>"
        )

        for group_name in group_names:

            rows = groups_data.get(
                group_name,
                [],
            )

            if index < len(
                rows
            ):

                row = rows[
                    index
                ]

                lines.append(
                    f"<td>{escape(row['school'])}</td>"
                )

                lines.append(
                    f"<td>{escape(row['project'])}</td>"
                )

                lines.append(
                    '<td align="center">'
                    f"{escape(row['score'])}"
                    "</td>"
                )

                lines.append(
                    '<td align="center">'
                    f"<strong>{escape(row['rank_text'])}</strong>"
                    "</td>"
                )

            else:

                lines.extend(
                    [
                        "<td></td>",
                        "<td></td>",
                        "<td></td>",
                        "<td></td>",
                    ]
                )

        lines.append(
            "</tr>"
        )

    lines.append(
        "</tbody>"
    )

    lines.append(
        "</table>"
    )

    lines.append("")

    lines.append(
        "> 项目、得分、排名及分组均以比赛实时接口为准；"
        "Excel 仅用于辅助补充学校名称。"
    )

    return "\n".join(
        lines
    )


# ============================================================
# 更新 README
# ============================================================

def update_readme(
    scoreboard_html,
):

    block = (
        f"{README_START}\n"
        f"{scoreboard_html}\n"
        f"{README_END}"
    )

    if README_FILE.exists():

        existing = README_FILE.read_text(
            encoding="utf-8"
        )

    else:

        existing = (
            "# 江苏大学生创新大赛实时成绩\n"
        )

    if (
        README_START in existing
        and README_END in existing
    ):

        pattern = re.compile(
            re.escape(
                README_START
            )
            + r".*?"
            + re.escape(
                README_END
            ),
            re.S,
        )

        new_content = pattern.sub(
            block,
            existing,
        )

    else:

        new_content = (
            existing.rstrip()
            + "\n\n"
            + block
            + "\n"
        )

    README_FILE.write_text(
        new_content,
        encoding="utf-8",
    )


# ============================================================
# 主程序
# ============================================================

def main():

    print()
    print("=" * 70)
    print("江苏大学生创新大赛实时成绩")
    print("=" * 70)

    print()

    print(
        "真实网页对应关系："
    )

    for group in GROUPS:

        print(
            f"  {group['name']} "
            f"→ road-show/{group['road_show_id']} "
            f"→ sort_type=2"
        )

    print()

    # ========================================================
    # Excel 仅补充学校
    # ========================================================

    (
        exact_lookup,
        normalized_lookup,
    ) = load_school_lookup()

    # ========================================================
    # API Session
    # ========================================================

    session = requests.Session()

    session.headers.update(
        HEADERS
    )

    # ========================================================
    # 获取五组数据
    # ========================================================

    groups_data = {}

    failed_groups = []

    for group in GROUPS:

        try:

            raw_data = fetch_group(
                session,
                group,
            )

            rows = parse_group(
                group,
                raw_data,
                exact_lookup,
                normalized_lookup,
            )

            groups_data[
                group["name"]
            ] = rows

        except Exception as exc:

            failed_groups.append(
                group["name"]
            )

            print(
                f"❌ {group['name']} 读取失败：{exc}",
                file=sys.stderr,
            )

        # 避免访问过快
        time.sleep(
            1
        )

    # ========================================================
    # 只要一组 API 失败，本轮都不覆盖旧榜单
    # ========================================================

    if failed_groups:

        print()

        print("=" * 70)

        print(
            "❌ 本次读取失败"
        )

        print(
            "失败组："
            + "、".join(
                failed_groups
            )
        )

        print(
            "为防止错误覆盖已有榜单，"
            "本次不会修改 README / CSV / JSON。"
        )

        print("=" * 70)

        sys.exit(
            1
        )

    # ========================================================
    # 输出统计
    # ========================================================

    total_projects = sum(
        len(
            rows
        )
        for rows in groups_data.values()
    )

    print("=" * 70)

    print(
        f"✅ 五个真实网页全部读取成功"
    )

    print(
        f"项目总数：{total_projects}"
    )

    print()

    for group in GROUPS:

        name = group[
            "name"
        ]

        print(
            f"{name} "
            f"(ID {group['road_show_id']})："
            f"{len(groups_data.get(name, []))} 项"
        )

    print("=" * 70)

    print()

    # ========================================================
    # 判断数据有没有变化
    # ========================================================

    current_data = comparable_data(
        groups_data
    )

    previous_data = load_previous_data()

    if previous_data == current_data:

        print(
            "✅ 当前数据与上一次一致。"
        )

        print(
            "无需修改 README / CSV / JSON。"
        )

        return

    # ========================================================
    # 数据发生变化
    # ========================================================

    now = datetime.now(
        ZoneInfo(
            "Asia/Shanghai"
        )
    ).strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    print(
        "🔄 检测到数据变化，开始更新。"
    )

    print(
        f"更新时间：{now}"
    )

    # ========================================================
    # JSON
    # ========================================================

    DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    json_payload = {
        "updated_at":
            now,

        "groups":
            current_data,
    }

    with JSON_FILE.open(
        "w",
        encoding="utf-8",
    ) as f:

        json.dump(
            json_payload,
            f,
            ensure_ascii=False,
            indent=2,
        )

    print(
        f"✅ {JSON_FILE}"
    )

    # ========================================================
    # CSV
    # ========================================================

    write_csv(
        groups_data
    )

    print(
        f"✅ {CSV_FILE}"
    )

    # ========================================================
    # README
    # ========================================================

    scoreboard_html = generate_scoreboard_html(
        groups_data,
        now,
    )

    update_readme(
        scoreboard_html
    )

    print(
        f"✅ {README_FILE}"
    )

    print()

    print("=" * 70)
    print(
        "✅ 本次更新完成"
    )
    print("=" * 70)


# ============================================================
# 程序入口
# ============================================================

if __name__ == "__main__":
    main()